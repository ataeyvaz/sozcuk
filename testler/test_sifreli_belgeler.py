"""Şifreli Word belgeleri (Dosya → Bilgi → Belgeyi Koru → Parolayla Şifrele) ve "Düzenlemeyi Kısıtla" testleri.

Çalıştırma (proje klasöründe):
    .venv\\Scripts\\python.exe -m unittest testler.test_sifreli_belgeler -v

Şifreli örnek dosyalar msoffcrypto-tool'un komut satırı aracıyla (-e -p) üretilir; Sözcük'ün kaydettiği dosyalar
da aynı araçla (-t, -p) denetlenir. Pencere testleri ekransız (offscreen) çalışır; kullanıcının kurtarma klasörüne
dokunulmaz (QStandardPaths test kipi).
"""

import io
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("SOZCUK_LOG", "0")

from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from PySide6.QtCore import QStandardPaths  # noqa: E402
from PySide6.QtGui import QTextDocument  # noqa: E402
from PySide6.QtWidgets import QApplication, QDialog  # noqa: E402

QStandardPaths.setTestModeEnabled(True)
APP = QApplication.instance() or QApplication([])

from sozcuk import docx_io, encryption, formats  # noqa: E402
from sozcuk.widgets import PasswordDialog  # noqa: E402

PASSWORD = "Gizli-Şifre 123"
WRONG = "yanlış"
TEXT = ["Şifreli belge denemesi: ğüşıöç ĞÜŞİÖÇ", "İkinci paragraf, kalın bir sözcük içerir."]


def cli(*args):
    """msoffcrypto-tool komut satırı aracı (pip ile gelen: python -m msoffcrypto)."""
    return subprocess.run([sys.executable, "-m", "msoffcrypto", *map(str, args)], capture_output=True)


def make_docx(path, protection=False):
    document = Document()
    document.add_paragraph(TEXT[0])
    paragraph = document.add_paragraph("İkinci paragraf, ")
    paragraph.add_run("kalın").bold = True
    paragraph.add_run(" bir sözcük içerir.")
    if protection:   # Word'ün "Düzenlemeyi Kısıtla → Yalnızca okuma" ayarı (şifreleme değil)
        settings = document.settings.element
        element = parse_xml(f'<w:documentProtection {nsdecls("w")} w:edit="readOnly" w:enforcement="1" '
                            'w:cryptProviderType="rsaAES" w:cryptAlgorithmClass="hash" w:cryptAlgorithmType="typeAny" '
                            'w:cryptAlgorithmSid="14" w:cryptSpinCount="100000" w:hash="abc=" w:salt="def="/>')
        settings.insert(1, element)
    document.save(path)


def paragraphs(document):
    return [block.text() for block in iter_blocks(document)]


def iter_blocks(document):
    block = document.begin()
    while block.isValid():
        yield block
        block = block.next()


def load(path, source=None):
    document = QTextDocument()
    docx_io.Reader(path, source).read(document)
    return document


class EncryptedFiles(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="sozcuk-test-")
        self.dir = Path(self.tmp.name)
        self.plain = self.dir / "duz.docx"
        self.encrypted = self.dir / "sifreli.docx"
        make_docx(self.plain)
        result = cli("-e", "-p", PASSWORD, self.plain, self.encrypted)
        self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))

    def tearDown(self):
        self.tmp.cleanup()

    # --- algılama ------------------------------------------------------

    def test_detection(self):
        self.assertEqual(encryption.probe(self.encrypted), encryption.ENCRYPTED_OOXML)
        self.assertIsNone(encryption.probe(self.plain))
        junk = self.dir / "bozuk.docx"
        junk.write_bytes(b"bu bir word belgesi degil")
        self.assertIsNone(encryption.probe(junk))
        # içerik OLE olduğu için sniff "ole" der ama .docx uzantısı docx olarak açılmaya devam eder
        fmt, problem = formats.effective_format(self.encrypted)
        self.assertEqual((fmt.kind, problem), ("docx", None))

    def test_plain_loader_gives_friendly_error(self):
        """Şifreli ya da bozuk dosya normal okuyucuya gelirse ham Python hatası (BadZipFile) değil Türkçe ileti."""
        with self.assertRaises(ValueError) as caught:
            docx_io.Reader(self.encrypted)
        self.assertIn("parolayla şifrelenmiş", str(caught.exception))
        junk = self.dir / "bozuk.docx"
        junk.write_bytes(b"PK\x03\x04 kirik zip")
        with self.assertRaises(ValueError) as caught:
            docx_io.Reader(junk)
        self.assertEqual(str(caught.exception), docx_io.UNREADABLE_PACKAGE)

    # --- açma ----------------------------------------------------------

    def test_correct_password_opens(self):
        before = set(self.dir.iterdir())
        package = encryption.decrypt(self.encrypted, PASSWORD)
        self.assertIsInstance(package, io.BytesIO)
        self.assertEqual(set(self.dir.iterdir()), before, "şifre çözülürken diske dosya yazılmamalı")
        self.assertEqual(paragraphs(load(self.encrypted, package))[:2], TEXT)

    def test_wrong_password_rejected(self):
        with self.assertRaises(encryption.WrongPassword):
            encryption.decrypt(self.encrypted, WRONG)
        with self.assertRaises(encryption.WrongPassword):
            encryption.decrypt(self.encrypted, PASSWORD.lower())   # büyük/küçük harf fark eder

    # --- kaydetme ------------------------------------------------------

    def test_round_trip_keeps_content_and_encryption(self):
        document = load(self.encrypted, encryption.decrypt(self.encrypted, PASSWORD))
        written = []
        real_replace = docx_io._replace_with_retry

        def spy(source, target):   # güvenli kayıttaki geçici dosya yalnızca şifreli bayt içermeli
            written.append(Path(source).read_bytes())
            real_replace(source, target)

        with mock.patch.object(docx_io, "_replace_with_retry", spy):
            docx_io.save(document, self.encrypted, PASSWORD)
        self.assertEqual(len(written), 1)
        self.assertTrue(written[0].startswith(formats.OLE_MAGIC), "geçici dosya şifreli (OLE) olmalı")
        self.assertNotIn(b"word/document.xml", written[0])
        self.assertEqual([p.name for p in self.dir.iterdir() if p.name.startswith("~$")], [])

        # msoffcrypto-tool CLI: dosya şifreli ve doğru parolayla çözülüyor
        self.assertEqual(cli("-t", self.encrypted).returncode, 0)
        self.assertIn(b"encrypted", cli("-t", "-v", self.encrypted).stdout + cli("-t", "-v", self.encrypted).stderr)
        decrypted = self.dir / "cozulmus.docx"
        result = cli("-p", PASSWORD, self.encrypted, decrypted)
        self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))
        self.assertTrue(zipfile.is_zipfile(decrypted))
        self.assertEqual(paragraphs(load(decrypted))[:2], TEXT)

        # Sözcük ile yeniden aç: içerik ve biçim (kalın) korunmuş
        reopened = load(self.encrypted, encryption.decrypt(self.encrypted, PASSWORD))
        self.assertEqual(paragraphs(reopened)[:2], TEXT)
        bold = [f.fragment().text() for f in iter_fragments(reopened) if f.fragment().charFormat().fontWeight() >= 600]
        self.assertIn("kalın", bold)
        with self.assertRaises(encryption.WrongPassword):
            encryption.decrypt(self.encrypted, WRONG)

    def test_verify(self):
        package = self.plain.read_bytes()
        data = encryption.encrypt(package, PASSWORD)
        self.assertTrue(encryption.verify(data, PASSWORD, package))
        self.assertFalse(encryption.verify(data, WRONG, package))
        self.assertFalse(encryption.verify(data, PASSWORD, package + b"x"))
        self.assertFalse(encryption.verify(b"bozuk", PASSWORD, package))

    def test_failed_verification_keeps_original(self):
        """Şifreli kayıt geri açılamıyorsa asıl dosyanın üzerine yazılmaz, geçici dosya kalmaz."""
        original = self.encrypted.read_bytes()
        document = load(self.encrypted, encryption.decrypt(self.encrypted, PASSWORD))
        real_encrypt = encryption.encrypt
        broken_cases = [
            lambda package, password: real_encrypt(package, password + "x"),     # başka parolayla
            lambda package, password: real_encrypt(self.plain.read_bytes(), password),     # başka içerik
            lambda package, password: real_encrypt(package, password)[:-4096],  # kesik dosya
        ]
        for broken in broken_cases:
            with mock.patch.object(encryption, "encrypt", broken):
                with self.assertRaises(encryption.EncryptedFileError) as caught:
                    docx_io.save(document, self.encrypted, PASSWORD)
            self.assertIsInstance(caught.exception, OSError, "kaydetme yolu OSError olarak yakalar")
            self.assertIn("üzerine yazılmadı", str(caught.exception))
            self.assertEqual(self.encrypted.read_bytes(), original, "özgün dosya değişmemeli")
            self.assertEqual([p.name for p in self.dir.iterdir() if p.name.startswith("~$")], [])

    def test_plain_save_without_password(self):
        document = load(self.encrypted, encryption.decrypt(self.encrypted, PASSWORD))
        target = self.dir / "sifresiz.docx"
        docx_io.save(document, target)
        self.assertTrue(zipfile.is_zipfile(target))
        self.assertIsNone(encryption.probe(target))

    def test_docm_from_memory(self):
        """Şifreli .docm de (içerik türü bellekte çevrilerek) bellekten açılır."""
        docm = self.dir / "makrolu.docm"
        encrypted_docm = self.dir / "makrolu-sifreli.docm"
        with zipfile.ZipFile(self.plain) as source, zipfile.ZipFile(docm, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(formats.OOXML_MAIN.encode(), formats.OOXML_VARIANTS[0].encode())
                target.writestr(item, data)
        self.assertEqual(cli("-e", "-p", PASSWORD, docm, encrypted_docm).returncode, 0)
        self.assertEqual(paragraphs(load(encrypted_docm, encryption.decrypt(encrypted_docm, PASSWORD)))[:2], TEXT)


def iter_fragments(document):
    for block in iter_blocks(document):
        iterator = block.begin()
        while not iterator.atEnd():
            yield iterator
            iterator += 1


class EditingRestriction(unittest.TestCase):
    """Word'ün "Düzenlemeyi Kısıtla"sı şifreleme değildir: parolasız açılmalı ve kayıtta korunmalı."""

    def test_opens_and_preserves_document_protection(self):
        with tempfile.TemporaryDirectory(prefix="sozcuk-test-") as tmp:
            source = Path(tmp) / "kisitli.docx"
            make_docx(source, protection=True)
            self.assertIsNone(encryption.probe(source))
            document = load(source)
            self.assertEqual(paragraphs(document)[:2], TEXT)
            self.assertTrue(docx_io.editing_restricted(document))

            for password in (None, PASSWORD):
                target = Path(tmp) / f"kayit-{bool(password)}.docx"
                docx_io.save(document, target, password)
                package = encryption.decrypt(target, password) if password else target
                settings = Document(package).settings.element
                protection = settings.find(qn("w:documentProtection"))
                self.assertIsNotNone(protection, "documentProtection kayıtta düşmemeli")
                self.assertEqual(protection.get(qn("w:edit")), "readOnly")
                self.assertEqual(protection.get(qn("w:enforcement")), "1")
                self.assertEqual(protection.get(qn("w:hash")), "abc=")
                # şema sırası: proofState'ten sonra, defaultTabStop'tan önce (Word sıraya katıdır)
                names = [child.tag.split("}")[1] for child in settings]
                index = names.index("documentProtection")
                self.assertLess(names.index("proofState"), index)
                self.assertLess(index, names.index("defaultTabStop"))

    def test_new_document_has_no_protection(self):
        document = QTextDocument()
        with tempfile.TemporaryDirectory(prefix="sozcuk-test-") as tmp:
            target = Path(tmp) / "yeni.docx"
            docx_io.save(document, target)
            self.assertIsNone(Document(target).settings.element.find(qn("w:documentProtection")))


class PasswordWindow(unittest.TestCase):
    """Parola penceresi ve ana penceredeki akış (ekransız)."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix="sozcuk-test-")
        cls.dir = Path(cls.tmp.name)
        cls.plain = cls.dir / "duz.docx"
        cls.encrypted = cls.dir / "sifreli.docx"
        make_docx(cls.plain)
        assert cli("-e", "-p", PASSWORD, cls.plain, cls.encrypted).returncode == 0

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def dialog(self):
        return PasswordDialog(message="deneme", accept_text="Aç",
                              verify=lambda password: encryption.decrypt(self.encrypted, password),
                              wrong_errors=(encryption.WrongPassword,))

    def test_wrong_then_correct(self):
        dialog = self.dialog()
        self.assertFalse(dialog.ok.isEnabled(), "boş parolayla Aç devre dışı")
        dialog.field.setText(WRONG)
        dialog._try_accept()
        self.assertNotEqual(dialog.result(), QDialog.Accepted)
        self.assertEqual(dialog.error.text(), "Şifre yanlış, tekrar deneyin.")
        self.assertFalse(dialog.error.isHidden())
        dialog.field.setText(PASSWORD)        # yeniden deneme: hata iletisi kalkar
        self.assertTrue(dialog.error.isHidden())
        dialog._try_accept()
        self.assertEqual(dialog.result(), QDialog.Accepted)
        self.assertEqual(dialog.password, PASSWORD)
        self.assertIsInstance(dialog.result_value, io.BytesIO)
        self.assertEqual(dialog.field.text(), "", "parola alanı kapanınca temizlenir")

    def test_show_hide_toggle(self):
        dialog = self.dialog()
        self.assertEqual(dialog.field.echoMode(), dialog.field.EchoMode.Password)
        dialog._toggle_visible()
        self.assertEqual(dialog.field.echoMode(), dialog.field.EchoMode.Normal)
        dialog._toggle_visible()
        self.assertEqual(dialog.field.echoMode(), dialog.field.EchoMode.Password)

    def test_confirm_mismatch(self):
        dialog = PasswordDialog(confirm=True)
        dialog.field.setText("bir")
        dialog.confirm_field.setText("iki")
        dialog._try_accept()
        self.assertNotEqual(dialog.result(), QDialog.Accepted)
        self.assertIn("eşleşmiyor", dialog.error.text())

    def test_main_window_flow(self):
        from sozcuk.window import MainWindow
        window = MainWindow()
        try:
            fmt = formats.BY_SUFFIX[".docx"]
            before = window.editor.document().toPlainText()

            # İptal: belge değişmez, parola tutulmaz
            with mock.patch.object(PasswordDialog, "exec", lambda self: QDialog.Rejected):
                window._open_encrypted(self.encrypted, fmt)
            self.assertEqual(window.editor.document().toPlainText(), before)
            self.assertIsNone(window.password)
            self.assertTrue(window.lock_indicator.isHidden())

            # Önce yanlış, sonra doğru parola
            def type_passwords(dialog):
                dialog.field.setText(WRONG)
                dialog._try_accept()
                assert dialog.error.text() == PasswordDialog.WRONG_PASSWORD
                dialog.field.setText(PASSWORD)
                dialog._try_accept()
                return dialog.result()

            with mock.patch.object(PasswordDialog, "exec", type_passwords):
                window._open_encrypted(self.encrypted, fmt)
            self.assertEqual(paragraphs(window.editor.document())[:2], TEXT)
            self.assertEqual(window.password, PASSWORD)
            self.assertFalse(window.lock_indicator.isHidden(), "durum çubuğunda Şifreli göstergesi")

            # Kaydet: aynı parolayla şifreli
            window.editor.textCursor().insertText("Eklenen ")
            self.assertTrue(window._write(window.path))
            self.assertEqual(encryption.probe(self.encrypted), encryption.ENCRYPTED_OOXML)
            self.assertTrue(paragraphs(load(self.encrypted, encryption.decrypt(self.encrypted, PASSWORD)))[0]
                            .startswith("Eklenen "))

            # Kurtarma kopyası da şifreli
            window.editor.textCursor().insertText("x")
            window._write_recovery_copy()
            recovery = window._recovery_file()
            self.assertEqual(encryption.probe(recovery), encryption.ENCRYPTED_OOXML)
            self.assertNotIn(PASSWORD, recovery.with_suffix(".json").read_text(encoding="utf-8"))
            window._discard_recovery()

            # Şifreyi kaldır → şifresiz kayıt, gösterge kalkar
            window.password = None
            self.assertTrue(window._write(window.path))
            self.assertIsNone(encryption.probe(self.encrypted))
            window._update_title()
            self.assertTrue(window.lock_indicator.isHidden())

            # Şifre ile koru (parola belirleme penceresi) → şifreli kayıt
            def set_password(dialog):
                dialog.field.setText(PASSWORD)
                dialog.confirm_field.setText(PASSWORD)
                dialog._try_accept()
                return dialog.result()

            with mock.patch.object(PasswordDialog, "exec", set_password):
                window.protect_with_password()
            self.assertEqual(window.password, PASSWORD)
            window._write(window.path)
            self.assertEqual(encryption.probe(self.encrypted), encryption.ENCRYPTED_OOXML)

            # Doğrulanamayan şifreli kayıt: kayıt başarısız sayılır, kullanıcıya söylenir, belge değişmiş kalır
            window.editor.textCursor().insertText("y")
            saved = self.encrypted.read_bytes()
            messages = []
            real_encrypt = encryption.encrypt
            with mock.patch.object(encryption, "encrypt", lambda p, pw: real_encrypt(p, pw + "x")), \
                    mock.patch("sozcuk.window.QMessageBox.critical", lambda *a: messages.append(a[2])):
                self.assertFalse(window._write(window.path))
            self.assertEqual(self.encrypted.read_bytes(), saved)
            self.assertTrue(window.editor.document().isModified())
            self.assertEqual(len(messages), 1)
            self.assertIn("doğrulanamadı", messages[0])
            self.assertNotIn("başka bir programda", messages[0])
            window._save_failed = False

            # Yeni belge: parola unutulur
            window.editor.document().setModified(False)
            window.new_document()
            self.assertIsNone(window.password)
        finally:
            window.editor.document().setModified(False)
            window._discard_recovery()
            window.close()


if __name__ == "__main__":
    unittest.main()
