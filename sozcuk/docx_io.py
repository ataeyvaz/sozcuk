"""Word (.docx) okuma ve yazma.

Desteklenenler: paragraflar, başlıklar, kalın/italik/altı çizili/üstü çizili, yazı tipi ve boyutu,
yazı ve vurgu rengi, üst/alt simge, hizalama, satır aralığı, paragraf boşlukları, girinti,
madde işaretli/numaralı listeler (iç içe), sayfa sonları, tablolar ve satır içi görseller.
Desteklenmeyenler (üst/alt bilgi, dipnot, yorum vb.) açılışta kullanıcıya bildirilir.
"""

import io
import os
import time
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Pt, RGBColor
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from lxml import etree
from PySide6.QtCore import QBuffer, QByteArray, QIODevice, Qt, QUrl
from PySide6.QtGui import (
    QBrush,
    QColor,
    QFont,
    QImage,
    QTextBlockFormat,
    QTextCharFormat,
    QTextDocument,
    QTextFormat,
    QTextListFormat,
    QTextTable,
)

from . import links, page_numbers, styles
from .builder import DocumentBuilder

EMU_PER_PX = 9525
PT_PER_PX = 72 / 96

ALIGN_TO_QT = {
    WD_ALIGN_PARAGRAPH.CENTER: Qt.AlignHCenter,
    WD_ALIGN_PARAGRAPH.RIGHT: Qt.AlignRight,
    WD_ALIGN_PARAGRAPH.JUSTIFY: Qt.AlignJustify,
    WD_ALIGN_PARAGRAPH.DISTRIBUTE: Qt.AlignJustify,
}

NUMFMT_TO_STYLE = {
    "decimal": QTextListFormat.ListDecimal,
    "decimalZero": QTextListFormat.ListDecimal,
    "lowerLetter": QTextListFormat.ListLowerAlpha,
    "upperLetter": QTextListFormat.ListUpperAlpha,
    "lowerRoman": QTextListFormat.ListLowerRoman,
    "upperRoman": QTextListFormat.ListUpperRoman,
}

STYLE_TO_NUMFMT = {
    QTextListFormat.ListDisc: ("bullet", "•"),
    QTextListFormat.ListCircle: ("bullet", "◦"),
    QTextListFormat.ListSquare: ("bullet", "▪"),
    QTextListFormat.ListDecimal: ("decimal", "%{n}."),
    QTextListFormat.ListLowerAlpha: ("lowerLetter", "%{n}."),
    QTextListFormat.ListUpperAlpha: ("upperLetter", "%{n}."),
    QTextListFormat.ListLowerRoman: ("lowerRoman", "%{n}."),
    QTextListFormat.ListUpperRoman: ("upperRoman", "%{n}."),
}

HIGHLIGHT_TO_HEX = {WD_COLOR_INDEX[name]: hex_color for name, hex_color in styles.HIGHLIGHT_COLORS}


# =============================================================================
# Okuma
# =============================================================================

DRAWING_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "v": "urn:schemas-microsoft-com:vml",
}


def _css_length_px(value):
    """VML stil uzunluğu (ör. "90pt", "3.2cm") -> piksel."""
    if not value:
        return None
    for unit, factor in (("pt", styles.PX_PER_PT), ("cm", styles.PX_PER_CM), ("mm", styles.PX_PER_CM / 10),
                         ("in", 96.0), ("px", 1.0)):
        if value.endswith(unit):
            try:
                return float(value[:-len(unit)]) * factor
            except ValueError:
                return None
    return None


class Reader:
    """Dosyayı kurucuda açar (hatalı dosya burada hata verir, mevcut belgeye dokunmadan);
    read() ile içeriği bir QTextDocument'e aktarır."""

    def __init__(self, path):
        from .formats import open_ooxml_package  # .docm/.dotx/.dotm da aynı paket biçimi
        self.docx = Document(open_ooxml_package(path))
        self.warnings = []
        self.numbering = self._numbering()
        self.theme_fonts = self._theme_fonts()
        self.defaults = self._doc_defaults()
        self.skipped = set()
        self.bookmarks = {}          # Word yer imi adı -> belge konumu (belge içi bağlantılar için)

    # --- belge geneli varsayılanlar --------------------------------------

    def _numbering(self):
        try:
            return self.docx.part.numbering_part.element
        except (NotImplementedError, KeyError, AttributeError):
            return None

    def _theme_fonts(self):
        fonts = {"minor": styles.DEFAULT_FAMILY, "major": styles.DEFAULT_FAMILY}
        for rel in self.docx.part.rels.values():
            if rel.reltype.endswith("/theme"):
                root = etree.fromstring(rel.target_part.blob)
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for key in ("minor", "major"):
                    found = root.xpath(f"//a:{key}Font/a:latin/@typeface", namespaces=ns)
                    if found and found[0]:
                        fonts[key] = found[0]
        return fonts

    def _fonts_from_rpr(self, rpr):
        """rPr içindeki w:rFonts'tan yazı tipi adını (tema fontları dahil) çözer."""
        if rpr is None:
            return None
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            return None
        name = rfonts.get(qn("w:ascii")) or rfonts.get(qn("w:hAnsi"))
        if name:
            return name
        theme = rfonts.get(qn("w:asciiTheme")) or rfonts.get(qn("w:hAnsiTheme"))
        if theme:
            return self.theme_fonts["major" if theme.startswith("major") else "minor"]
        return None

    def _doc_defaults(self):
        defaults = {"family": self.theme_fonts["minor"], "size": 10.0,
                    "after": 0.0, "before": 0.0, "line": 1.0}
        root = self.docx.styles.element
        rpr = root.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
        if rpr is not None:
            defaults["family"] = self._fonts_from_rpr(rpr) or defaults["family"]
            sz = rpr.find(qn("w:sz"))
            if sz is not None:
                defaults["size"] = int(sz.get(qn("w:val"))) / 2
        spacing = root.find(f"{qn('w:docDefaults')}/{qn('w:pPrDefault')}/{qn('w:pPr')}/{qn('w:spacing')}")
        if spacing is not None:
            if spacing.get(qn("w:after")):
                defaults["after"] = int(spacing.get(qn("w:after"))) / 20
            if spacing.get(qn("w:before")):
                defaults["before"] = int(spacing.get(qn("w:before"))) / 20
            if spacing.get(qn("w:line")) and spacing.get(qn("w:lineRule"), "auto") == "auto":
                defaults["line"] = int(spacing.get(qn("w:line"))) / 240
        return defaults

    # --- stil zinciri ----------------------------------------------------

    @staticmethod
    def _chain(style):
        while style is not None:
            yield style
            style = style.base_style

    def _style_font(self, style, attr):
        for s in self._chain(style):
            if attr == "name":
                name = self._fonts_from_rpr(s.element.rPr)
                if name:
                    return name
                continue
            value = getattr(s.font, attr)
            if value is not None:
                return value
        return None

    def _para_value(self, par, attr):
        value = getattr(par.paragraph_format, attr)
        if value is not None:
            return value
        for s in self._chain(par.style):
            value = getattr(s.paragraph_format, attr)
            if value is not None:
                return value
        return None

    @staticmethod
    def _heading_level(style):
        for s in Reader._chain(style):
            name = (s.name or "").lower()
            if name in ("title", "heading 1"):
                return 1
            if name.startswith("heading ") and name[8:].isdigit():
                return 2
        return 0

    # --- listeler --------------------------------------------------------

    def _num_pr(self, par):
        candidates = [par._p.pPr] + [s.element.pPr for s in self._chain(par.style)]
        for ppr in candidates:
            if ppr is not None and ppr.numPr is not None:
                num_pr = ppr.numPr
                num_id = num_pr.numId.val if num_pr.numId is not None else None
                ilvl = num_pr.ilvl.val if num_pr.ilvl is not None else 0
                if num_id is not None:
                    return num_id, ilvl
        return None

    def _list_style(self, num_id, ilvl):
        fmt = "bullet"
        if self.numbering is not None:
            abstract = self.numbering.xpath(f'w:num[@w:numId="{num_id}"]/w:abstractNumId/@w:val')
            if abstract:
                found = self.numbering.xpath(
                    f'w:abstractNum[@w:abstractNumId="{abstract[0]}"]/w:lvl[@w:ilvl="{ilvl}"]/w:numFmt/@w:val'
                )
                if found:
                    fmt = found[0]
        if fmt == "none":
            return None
        if fmt == "bullet":
            return styles.list_style_for(True, ilvl + 1)
        return NUMFMT_TO_STYLE.get(fmt, QTextListFormat.ListDecimal)

    # --- içerik ----------------------------------------------------------

    def read(self, document):
        self.warnings = []
        self.skipped = set()
        self.qdoc = document
        recognized = self._read_page_numbers(document)
        self._check_unsupported(recognized)
        self._read_page_setup(document)
        self._blocks(DocumentBuilder(document), self.docx.iter_inner_content())
        links.set_bookmarks(document, self.bookmarks)
        if self.skipped:
            self.warnings.append("Desteklenmeyen öğeler atlandı: " + ", ".join(sorted(self.skipped)) + ".")
        return self.warnings

    def _read_page_setup(self, document):
        """İlk bölümün kâğıt boyutu, yönlendirmesi ve kenar boşlukları (Word'de de bölüm düzeyindedir)."""
        if not self.docx.sections:
            return
        section = self.docx.sections[0]

        def px(length):
            return None if length is None else length.pt / PT_PER_PX

        width, height = px(section.page_width), px(section.page_height)
        if width and height and section.orientation == WD_ORIENT.LANDSCAPE and width < height:
            width, height = height, width  # yönlendirmesi yatay ama ölçüleri dikey yazılmış dosyalar
        styles.set_page_setup(document, width=width, height=height,
                              top=px(section.top_margin), bottom=px(section.bottom_margin),
                              left=px(section.left_margin), right=px(section.right_margin))

    def _read_page_numbers(self, document):
        """Word'ün sayfa numarası alanını (üst/alt bilgideki PAGE / NUMPAGES) Sözcük ayarına çevirir.
        Tanınan üst/alt bilgi bölümlerini döndürür (bunlar "desteklenmiyor" diye bildirilmez)."""
        recognized = set()
        for index, section in enumerate(self.docx.sections):
            for side, part in (("top", section.header), ("bottom", section.footer)):
                if part.is_linked_to_previous:
                    continue
                for paragraph in part.paragraphs:
                    field = self._page_number_field(paragraph)
                    if field is None:
                        continue
                    has_total, literal, tabs = field
                    # hizalama sırasıyla: paragrafın kendi hizası, Word'ün sayfa numarası çerçevesi
                    # (w:framePr w:xAlign), yoksa alt bilgi stilindeki sekmeler (1 sekme orta, 2 sekme sağ)
                    frame = paragraph._p.xpath("./w:pPr/w:framePr/@w:xAlign")
                    horizontal = {WD_ALIGN_PARAGRAPH.CENTER: "center", WD_ALIGN_PARAGRAPH.RIGHT: "right"}.get(
                        paragraph.alignment,
                        frame[0] if frame and frame[0] in ("center", "right", "left")
                        else "center" if tabs == 1 else "right" if tabs >= 2 else "left")
                    position = f"{side}_{horizontal}"
                    fmt = "x_of_y" if has_total else ("page_x" if any(ch.isalpha() for ch in literal) else "plain")
                    page_numbers.set_settings(document, enabled=True, format=fmt, start=self._page_number_start(section),
                                              position=position if position in page_numbers.POSITION_BY_KEY
                                              else page_numbers.DEFAULT_POSITION)
                    recognized.add((index, side))
                    break
        return recognized

    @staticmethod
    def _page_number_field(paragraph):
        """Paragrafta sayfa numarası alanı varsa (NUMPAGES var mı, alan dışındaki düz metin, alandan önceki
        sekme sayısı) döndürür; yoksa None."""
        codes = " ".join((node.text or "") for node in paragraph._p.xpath(".//w:instrText")).upper()
        if "PAGE" not in codes:
            return None
        depth, literal, tabs, started = 0, [], 0, False
        for run in paragraph._p.xpath("./w:r"):
            for child in run:
                tag = child.tag
                if tag == qn("w:fldChar"):
                    kind = child.get(qn("w:fldCharType"))
                    if kind == "begin":
                        started = True
                    depth += 1 if kind == "begin" else -1 if kind == "end" else 0
                elif tag == qn("w:tab") and not started:
                    tabs += 1
                elif tag == qn("w:t") and depth <= 0:
                    literal.append(child.text or "")
        return ("NUMPAGES" in codes, "".join(literal), tabs)

    @staticmethod
    def _page_number_start(section):
        found = section._sectPr.xpath("./w:pgNumType/@w:start")
        try:
            return max(0, int(found[0])) if found else 1
        except (TypeError, ValueError):
            return 1

    def _check_unsupported(self, recognized=()):
        for index, section in enumerate(self.docx.sections):
            for side, part_name, part in (("top", "üst bilgi", section.header), ("bottom", "alt bilgi", section.footer)):
                if (index, side) in recognized:
                    continue   # yalnızca sayfa numarası alanı: Sözcük bunu destekliyor
                if not part.is_linked_to_previous and any(p.text.strip() for p in part.paragraphs):
                    self.skipped.add(part_name)
        body = self.docx.element.body
        if body.xpath(".//w:footnoteReference"):
            self.skipped.add("dipnotlar")
        if body.xpath(".//w:commentRangeStart"):
            self.skipped.add("yorumlar")
        if body.xpath(".//w:txbxContent"):
            self.skipped.add("metin kutuları")

    def _blocks(self, builder, items):
        for item in items:
            if isinstance(item, Paragraph):
                self._paragraph(builder, item)
            elif isinstance(item, Table):
                self._table(builder, item)

    def _paragraph(self, builder, par):
        level = self._heading_level(par.style)
        block = QTextBlockFormat()
        block.setHeadingLevel(level)

        alignment = self._para_value(par, "alignment")
        block.setAlignment(ALIGN_TO_QT.get(alignment, Qt.AlignLeft))

        before = self._para_value(par, "space_before")
        after = self._para_value(par, "space_after")
        block.setTopMargin((before.pt if before is not None else self.defaults["before"]) / PT_PER_PX)
        block.setBottomMargin((after.pt if after is not None else self.defaults["after"]) / PT_PER_PX)

        line = self._para_value(par, "line_spacing")
        if line is None:
            line = self.defaults["line"]
        if isinstance(line, float):
            block.setLineHeight(line * 100, QTextBlockFormat.ProportionalHeight.value)
        elif line is not None:
            block.setLineHeight(line.pt / PT_PER_PX, QTextBlockFormat.MinimumHeight.value)

        num = self._num_pr(par)
        list_style = self._list_style(*num) if num and num[0] != 0 else None
        if list_style is None:
            left = self._para_value(par, "left_indent")
            if left:
                block.setLeftMargin(left.pt / PT_PER_PX)
            first = self._para_value(par, "first_line_indent")
            if first:
                block.setTextIndent(first.pt / PT_PER_PX)
        right = self._para_value(par, "right_indent")
        if right:
            block.setRightMargin(right.pt / PT_PER_PX)

        if self._para_value(par, "page_break_before"):
            block.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)

        base = self._base_char_format(par.style)
        builder.paragraph(block, base)
        if list_style is not None:
            num_id, ilvl = num
            builder.list_item((num_id, ilvl), list_style, ilvl + 1)

        for name in par._p.xpath("./w:bookmarkStart/@w:name"):
            self.bookmarks.setdefault(name, builder.cursor.position())

        for item in par.iter_inner_content():
            if isinstance(item, Hyperlink):
                href = self._hyperlink_href(item)
                runs = item.runs
            else:
                href, runs = None, [item]
            for run in runs:
                self._run(builder, run, base, href)

    def _base_char_format(self, style):
        fmt = QTextCharFormat()
        fmt.setFontFamilies([self._style_font(style, "name") or self.defaults["family"]])
        size = self._style_font(style, "size")
        fmt.setFontPointSize(size.pt if size is not None else self.defaults["size"])
        fmt.setFontWeight(QFont.Bold if self._style_font(style, "bold") else QFont.Normal)
        fmt.setFontItalic(bool(self._style_font(style, "italic")))
        if self._style_font(style, "underline"):
            fmt.setFontUnderline(True)
        if self._style_font(style, "strike"):
            fmt.setFontStrikeOut(True)
        for s in self._chain(style):
            if s.font.color is not None and s.font.color.rgb is not None:
                fmt.setForeground(QBrush(QColor(f"#{s.font.color.rgb}")))
                break
        return fmt

    @staticmethod
    def _hyperlink_href(hyperlink):
        """Word köprüsünü Sözcük adresine çevirir (dış adres, Sözcük'ün sayfa/satır yer imi ya da yer imi adı)."""
        address = hyperlink.address or ""
        fragment = hyperlink.fragment or ""
        if address:
            return f"{address}#{fragment}" if fragment else address
        if fragment:
            match = links.DOCX_BOOKMARK.match(fragment)
            if match:
                return links.page_target(int(match.group(1)), int(match.group(2)))
            return links.bookmark_target(fragment)
        return None

    def _run(self, builder, run, base, href=None):
        fmt = QTextCharFormat(base)
        if run.style is not None and run.style.name != "Default Paragraph Font":
            for attr in ("bold", "italic", "underline", "strike"):
                if self._style_font(run.style, attr) is not None:
                    self._apply_font_attr(fmt, attr, self._style_font(run.style, attr))
        font = run.font
        name = self._fonts_from_rpr(run._r.rPr)
        if name:
            fmt.setFontFamilies([name])
        if font.size is not None:
            fmt.setFontPointSize(font.size.pt)
        for attr in ("bold", "italic", "underline", "strike"):
            value = getattr(font, attr)
            if value is not None:
                self._apply_font_attr(fmt, attr, value)
        if font.color is not None and font.color.type is not None and font.color.rgb is not None:
            fmt.setForeground(QBrush(QColor(f"#{font.color.rgb}")))
        if font.highlight_color in HIGHLIGHT_TO_HEX:
            fmt.setBackground(QBrush(QColor(HIGHLIGHT_TO_HEX[font.highlight_color])))
        if font.superscript:
            fmt.setVerticalAlignment(QTextCharFormat.AlignSuperScript)
        elif font.subscript:
            fmt.setVerticalAlignment(QTextCharFormat.AlignSubScript)
        if href:
            fmt = links.link_format(fmt)
            fmt.setAnchorHref(href)

        for child in run._r.iterchildren():
            tag = child.tag
            if tag == qn("w:t"):
                builder.text(child.text or "", fmt)
            elif tag == qn("w:tab"):
                builder.text("\t", fmt)
            elif tag in (qn("w:br"), qn("w:cr")):
                if child.get(qn("w:type")) == "page":
                    block_fmt = builder.cursor.blockFormat()
                    block_fmt.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysAfter)
                    builder.cursor.setBlockFormat(block_fmt)
                else:
                    builder.text("\u2028", fmt)
            elif tag in (qn("w:drawing"), qn("w:pict")):
                self._drawing(builder, run, child)
            elif tag == qn("w:noBreakHyphen"):
                builder.text("\u2011", fmt)

    @staticmethod
    def _apply_font_attr(fmt, attr, value):
        on = bool(value)
        if attr == "bold":
            fmt.setFontWeight(QFont.Bold if on else QFont.Normal)
        elif attr == "italic":
            fmt.setFontItalic(on)
        elif attr == "underline":
            fmt.setFontUnderline(on)
        elif attr == "strike":
            fmt.setFontStrikeOut(on)

    def _drawing(self, builder, run, element):
        # eski Word belgelerinden dönüştürülen dosyalarda resim VML (w:pict / v:imagedata) olarak gelir
        blips = etree.ElementBase.xpath(element, ".//a:blip/@r:embed | .//v:imagedata/@r:id", namespaces=DRAWING_NS)
        if not blips:
            self.skipped.add("şekiller ve çizimler")
            return
        try:
            blob = run.part.related_parts[blips[0]].blob
        except KeyError:
            return
        image = QImage.fromData(blob)
        if image.isNull():
            self.skipped.add("desteklenmeyen biçimdeki görseller")
            return
        width, height = image.width(), image.height()
        extent = etree.ElementBase.xpath(element, ".//wp:extent", namespaces=DRAWING_NS)
        shape = etree.ElementBase.xpath(element, ".//v:shape/@style", namespaces=DRAWING_NS)
        if extent:
            width = int(extent[0].get("cx")) / EMU_PER_PX
            height = int(extent[0].get("cy")) / EMU_PER_PX
        elif shape:
            size = dict(part.split(":", 1) for part in shape[0].replace(" ", "").split(";") if ":" in part)
            width = _css_length_px(size.get("width")) or width
            height = _css_length_px(size.get("height")) or height
        max_width = styles.text_width(self.qdoc)
        if width > max_width:
            height, width = height * max_width / width, max_width
        if etree.ElementBase.xpath(element, ".//wp:anchor", namespaces=DRAWING_NS):
            self.skipped.add("görsellerin serbest konumu (satır içine alındı)")
        builder.image(image, width, height)

    def _table(self, builder, table):
        rows, cols = len(table.rows), len(table.columns)
        if not rows or not cols:
            return
        qtable = builder.table(rows, cols)
        seen = []  # birleştirilmiş hücreler aynı w:tc öğesini döndürür
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells[:cols]):
                if any(cell._tc is tc for tc in seen):
                    continue
                seen.append(cell._tc)
                self._blocks(builder.cell(qtable, r, c), cell.iter_inner_content())


def load(path, document):
    """docx dosyasını document'e yükler; kullanıcıya gösterilecek uyarıları döndürür."""
    return Reader(path).read(document)


# =============================================================================
# Yazma
# =============================================================================

class _Writer:
    def __init__(self, qdoc):
        self.qdoc = qdoc
        self.docx = Document()
        self.block_paragraphs = {}   # Qt blok konumu -> docx paragrafı (yer imleri için)
        self.bookmark_targets = {}   # yer imi adı -> (sayfa, satır)
        self.list_nums = {}       # QTextList.objectIndex() -> numId
        self.abstract_ids = {}    # (numFmt, metin) -> abstractNumId
        self._setup()

    def _setup(self):
        section = self.docx.sections[0]
        width, height = styles.page_size(self.qdoc)
        section.orientation = WD_ORIENT.LANDSCAPE if width > height else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Emu(round(width * EMU_PER_PX)), Emu(round(height * EMU_PER_PX))
        for side, value in zip(("top_margin", "bottom_margin", "left_margin", "right_margin"),
                               styles.all_margins(self.qdoc)):
            setattr(section, side, Emu(round(value * EMU_PER_PX)))

        self._write_page_numbers(section)

        normal = self.docx.styles["Normal"]
        normal.font.name = styles.DEFAULT_FAMILY
        normal.font.size = Pt(styles.DEFAULT_SIZE)
        for level in (1, 2):
            style = self.docx.styles[f"Heading {level}"]
            size, bold, before, after = styles.HEADING_STYLES[level]
            style.font.name = styles.DEFAULT_FAMILY
            style.font.size = Pt(size)
            style.font.bold = bold
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)

        body = self.docx.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)

    def write(self, path):
        self._frame(self.qdoc.rootFrame(), self.docx)
        self._write_bookmarks()
        if not self.docx.paragraphs and not self.docx.tables:
            self.docx.add_paragraph()
        path = Path(path)
        # Önce aynı klasörde geçici dosyaya yazılıp tek adımda yerine konur: eşitleme uygulaması (Google Drive,
        # OneDrive) yarım yazılmış belgeyi asla görmez. "~$" ile başlayan adları Office gibi eşitleme
        # uygulamaları da geçici dosya sayıp buluta göndermez.
        tmp = path.with_name(f"~$sozcuk-{os.getpid()}-{path.name}")
        try:
            self.docx.save(tmp)
            _replace_with_retry(tmp, path)
        except BaseException:
            try:
                tmp.unlink(missing_ok=True)
            except OSError:
                pass
            raise

    def _write_page_numbers(self, section):
        """Sayfa numarasını Word'ün kendi alanı olarak yazar: dosya Word'de açıldığında canlı PAGE / NUMPAGES
        alanı olarak kalır ve sayfa eklenince kendiliğinden güncellenir."""
        config = page_numbers.settings(self.qdoc)
        if not config.enabled:
            return
        part = section.header if config.side == "top" else section.footer
        part.is_linked_to_previous = False
        paragraph = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
        for element in list(paragraph._p):
            if element.tag != qn("w:pPr"):
                paragraph._p.remove(element)
        paragraph.alignment = {Qt.AlignHCenter: WD_ALIGN_PARAGRAPH.CENTER,
                               Qt.AlignRight: WD_ALIGN_PARAGRAPH.RIGHT}.get(config.alignment,
                                                                            WD_ALIGN_PARAGRAPH.LEFT)
        for text, field in _page_number_parts(config.pattern):
            if field:
                self._field_runs(paragraph, field)
            else:
                run = paragraph.add_run(text)
                run.font.name = styles.DEFAULT_FAMILY
                run.font.size = Pt(page_numbers.FONT_SIZE)
        if config.start != 1:
            pg_num = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="{config.start}"/>')
            section._sectPr.insert_element_before(
                pg_num, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection",
                "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")

    def _field_runs(self, paragraph, code):
        """{ PAGE } / { NUMPAGES } alanını oluşturan run dizisi (Word'ün yazdığı biçimde)."""
        for xml in (f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>',
                    f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="{styles.DEFAULT_FAMILY}" '
                    f'w:hAnsi="{styles.DEFAULT_FAMILY}"/><w:sz w:val="{int(page_numbers.FONT_SIZE * 2)}"/></w:rPr>'
                    f'<w:instrText xml:space="preserve"> {code} </w:instrText></w:r>',
                    f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>',
                    f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="{styles.DEFAULT_FAMILY}" '
                    f'w:hAnsi="{styles.DEFAULT_FAMILY}"/><w:sz w:val="{int(page_numbers.FONT_SIZE * 2)}"/></w:rPr>'
                    f'<w:t>1</w:t></w:r>',
                    f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'):
            paragraph._p.append(parse_xml(xml))

    def _frame(self, frame, container):
        for it in frame:
            child = it.currentFrame()
            if child is not None:
                if isinstance(child, QTextTable):
                    self._table(child, container)
                else:
                    self._frame(child, container)
                continue
            block = it.currentBlock()
            if block.isValid():
                self._block(block, container)

    def _table(self, table, container):
        rows, cols = table.rows(), table.columns()
        wtable = container.add_table(rows=rows, cols=cols)
        wtable.style = self.docx.styles["Table Grid"]
        for r in range(rows):
            for c in range(cols):
                cell = table.cellAt(r, c)
                if cell.row() != r or cell.column() != c:
                    continue
                wcell = wtable.cell(r, c)
                if cell.rowSpan() > 1 or cell.columnSpan() > 1:
                    wcell = wcell.merge(wtable.cell(r + cell.rowSpan() - 1, c + cell.columnSpan() - 1))
                first = True
                block = cell.firstCursorPosition().block()
                last = cell.lastCursorPosition().block()
                while block.isValid():
                    self._block(block, wcell, reuse=first)
                    first = False
                    if block == last:
                        break
                    block = block.next()

    def _block(self, block, container, reuse=False):
        par = container.paragraphs[0] if reuse and container.paragraphs else container.add_paragraph()
        self.block_paragraphs[block.position()] = par
        fmt = block.blockFormat()
        pf = par.paragraph_format

        level = fmt.headingLevel()
        if level:
            par.style = self.docx.styles[f"Heading {min(level, 2)}"]

        lst = block.textList()
        if lst is not None:
            self._list(par, lst)
        else:
            left = fmt.indent() * styles.INDENT_WIDTH + fmt.leftMargin()
            if left:
                pf.left_indent = Emu(round(left * EMU_PER_PX))
        if fmt.textIndent():
            pf.first_line_indent = Emu(round(fmt.textIndent() * EMU_PER_PX))
        if fmt.rightMargin():
            pf.right_indent = Emu(round(fmt.rightMargin() * EMU_PER_PX))

        align = fmt.alignment() & Qt.AlignHorizontal_Mask
        if align & Qt.AlignHCenter:
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align & Qt.AlignJustify:
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align & (Qt.AlignRight | Qt.AlignTrailing) and not align & Qt.AlignLeft:
            pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        pf.space_before = Pt(round(fmt.topMargin() * PT_PER_PX, 1))
        pf.space_after = Pt(round(fmt.bottomMargin() * PT_PER_PX, 1))
        if fmt.lineHeightType() == QTextBlockFormat.ProportionalHeight.value:
            pf.line_spacing = fmt.lineHeight() / 100
        elif fmt.lineHeightType() in (QTextBlockFormat.MinimumHeight.value, QTextBlockFormat.FixedHeight.value):
            pf.line_spacing = Pt(fmt.lineHeight() * PT_PER_PX)
        if fmt.pageBreakPolicy() & QTextFormat.PageBreak_AlwaysBefore:
            pf.page_break_before = True

        for it in block:
            fragment = it.fragment()
            if not fragment.isValid():
                continue
            cf = fragment.charFormat()
            if cf.isImageFormat():
                self._image(par, cf.toImageFormat())
                continue
            text = fragment.text().replace("\u2028", "\n").replace("\ufffc", "")
            if text:
                run = par.add_run(text)
                self._char(run, cf, bool(level))
                if cf.anchorHref():
                    self._hyperlink(par, run, cf.anchorHref())

        if fmt.pageBreakPolicy() & QTextFormat.PageBreak_AlwaysAfter:
            par.add_run().add_break(WD_BREAK.PAGE)

    def _hyperlink(self, par, run, href):
        """Metni Word'ün köprü öğesine sarar: web adresleri ilişkiyle, belge içi hedefler yer imiyle."""
        kind, value = links.parse(href)
        if kind == "web":
            rel_id = par.part.relate_to(value, RT.HYPERLINK, is_external=True)
            element = parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rel_id}"/>')
        elif kind in ("page", "bookmark"):
            if kind == "page":
                page, line = value
                name = links.bookmark_name(page, line)
                self.bookmark_targets[name] = (page, line)
            else:
                name = value
            element = parse_xml(f'<w:hyperlink {nsdecls("w")} w:anchor="{name}"/>')
        else:
            return
        run._element.addprevious(element)
        element.append(run._element)

    def _write_bookmarks(self):
        """Belge içi bağlantıların hedeflerine yer imi koyar (Word'de de çalışsın diye)."""
        for index, (name, (page, line)) in enumerate(self.bookmark_targets.items(), start=1):
            position = links.position_for(self.qdoc, page, line)
            par = self.block_paragraphs.get(self.qdoc.findBlock(position).position())
            if par is None:
                continue
            start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{index}" w:name="{name}"/>')
            end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{index}"/>')
            par._p.insert(1 if par._p.find(qn("w:pPr")) is not None else 0, start)
            par._p.append(end)

    def _char(self, run, cf, heading):
        font = run.font
        families = cf.fontFamilies()
        if families:
            font.name = families[0]
        if cf.fontPointSize() > 0:
            font.size = Pt(cf.fontPointSize())
        bold = cf.fontWeight() >= 600
        if bold or heading:
            font.bold = bold
        if cf.fontItalic():
            font.italic = True
        if cf.fontUnderline():
            font.underline = True
        if cf.fontStrikeOut():
            font.strike = True
        fg = cf.foreground()
        if cf.hasProperty(QTextFormat.ForegroundBrush) and fg.style() != Qt.NoBrush:
            c = fg.color()
            font.color.rgb = RGBColor(c.red(), c.green(), c.blue())
        bg = cf.background()
        if cf.hasProperty(QTextFormat.BackgroundBrush) and bg.style() != Qt.NoBrush and bg.color().alpha():
            font.highlight_color = self._nearest_highlight(bg.color())
        valign = cf.verticalAlignment()
        if valign == QTextCharFormat.AlignSuperScript:
            font.superscript = True
        elif valign == QTextCharFormat.AlignSubScript:
            font.subscript = True

    @staticmethod
    def _nearest_highlight(color):
        def distance(item):
            c = QColor(item[1])
            return (c.red() - color.red()) ** 2 + (c.green() - color.green()) ** 2 + (c.blue() - color.blue()) ** 2
        name, _ = min(styles.HIGHLIGHT_COLORS, key=distance)
        return WD_COLOR_INDEX[name]

    def _image(self, par, image_format):
        resource = self.qdoc.resource(QTextDocument.ImageResource, QUrl(image_format.name()))
        image = resource if isinstance(resource, QImage) else QImage(resource) if resource is not None else QImage()
        if image.isNull():
            return
        data = QByteArray()
        buffer = QBuffer(data)
        buffer.open(QIODevice.WriteOnly)
        image.save(buffer, "PNG")
        buffer.close()
        width = image_format.width() or image.width()
        height = image_format.height() or image.height()
        par.add_run().add_picture(
            io.BytesIO(bytes(data)),
            width=Emu(round(width * EMU_PER_PX)),
            height=Emu(round(height * EMU_PER_PX)),
        )

    # --- listeler ----------------------------------------------------------

    def _list(self, par, lst):
        fmt = lst.format()
        ilvl = max(0, min(8, fmt.indent() - 1))
        key = lst.objectIndex()  # Python sarmalayıcısı her çağrıda yenidir, id() kullanılamaz
        if key not in self.list_nums:
            num_fmt, text = STYLE_TO_NUMFMT.get(fmt.style(), ("decimal", "%{n}."))
            abstract_id = self._abstract(num_fmt, text)
            self.list_nums[key] = self._num(abstract_id)
        par.style = self.docx.styles["List Paragraph"]
        num_pr = par._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = ilvl
        num_pr.get_or_add_numId().val = self.list_nums[key]

    def _numbering(self):
        return self.docx.part.numbering_part.element

    def _abstract(self, num_fmt, text):
        if (num_fmt, text) in self.abstract_ids:
            return self.abstract_ids[(num_fmt, text)]
        numbering = self._numbering()
        existing = [int(v) for v in numbering.xpath("w:abstractNum/@w:abstractNumId")]
        abstract_id = max(existing, default=-1) + 1
        levels = []
        for i in range(9):
            lvl_text = text.replace("{n}", str(i + 1))
            left = 720 * (i + 1)
            levels.append(
                f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{num_fmt}"/>'
                f'<w:lvlText w:val="{lvl_text}"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="{left}" w:hanging="360"/></w:pPr></w:lvl>'
            )
        element = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
            f'<w:multiLevelType w:val="hybridMultilevel"/>{"".join(levels)}</w:abstractNum>'
        )
        nums = numbering.xpath("w:num")
        if nums:
            nums[0].addprevious(element)
        else:
            numbering.append(element)
        self.abstract_ids[(num_fmt, text)] = abstract_id
        return abstract_id

    def _num(self, abstract_id):
        numbering = self._numbering()
        existing = [int(v) for v in numbering.xpath("w:num/@w:numId")]
        num_id = max(existing, default=0) + 1
        overrides = "".join(
            f'<w:lvlOverride w:ilvl="{i}"><w:startOverride w:val="1"/></w:lvlOverride>' for i in range(9)
        )
        numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{overrides}</w:num>'
        ))
        return num_id


REPLACE_ATTEMPTS = 8
REPLACE_DELAY = 0.25


def _replace_with_retry(source, target):
    """Eşitleme uygulaması hedef dosyayı buluta gönderirken kısa süre kilitleyebilir; birkaç kez yeniden dene."""
    for attempt in range(REPLACE_ATTEMPTS):
        try:
            os.replace(source, target)
            return
        except PermissionError:
            if attempt == REPLACE_ATTEMPTS - 1:
                raise
            time.sleep(REPLACE_DELAY)


def _page_number_parts(pattern):
    """Deseni ("Sayfa {page} / {total}") düz metin ve alan parçalarına ayırır."""
    parts, text = [], ""
    index = 0
    while index < len(pattern):
        for token, code in (("{page}", "PAGE"), ("{total}", "NUMPAGES")):
            if pattern.startswith(token, index):
                if text:
                    parts.append((text, None))
                    text = ""
                parts.append((None, code))
                index += len(token)
                break
        else:
            text += pattern[index]
            index += 1
    if text:
        parts.append((text, None))
    return parts


def save(document, path):
    _Writer(document).write(path)
